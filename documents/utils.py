import io
import os
import re
import docx
import pdf2image
import pdfplumber
import pytesseract
from .models import Document


def extract_text(file_obj_or_path):
    """
    Extract text content from a PDF or DOCX file.
    Accepts a filepath, an UploadedFile, or a file-like object.
    Returns extracted text string, or empty string if extraction fails or finds no text.
    If pdfplumber finds no text in a PDF, it falls back to OCR via pdf2image and pytesseract (up to 20 pages).
    """
    if not file_obj_or_path:
        return ''

    # Determine file name to detect format
    filename = ''
    if hasattr(file_obj_or_path, 'name'):
        filename = file_obj_or_path.name
    elif isinstance(file_obj_or_path, str):
        filename = file_obj_or_path

    ext = os.path.splitext(filename)[1].lower()

    text_parts = []

    try:
        if ext == '.pdf':
            pdf_bytes = None
            # Handle UploadedFile or file-like or path
            if hasattr(file_obj_or_path, 'read'):
                file_obj_or_path.seek(0)
                pdf_bytes = file_obj_or_path.read()
                file_obj_or_path.seek(0)
                with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(page_text)
            else:
                with pdfplumber.open(file_obj_or_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(page_text)

            # If pdfplumber finds no text at all, fall back to OCR
            if not any(part.strip() for part in text_parts):
                try:
                    if pdf_bytes is not None:
                        images = pdf2image.convert_from_bytes(pdf_bytes, first_page=1, last_page=20)
                    else:
                        images = pdf2image.convert_from_path(file_obj_or_path, first_page=1, last_page=20)

                    ocr_parts = []
                    for img in images[:20]:
                        page_ocr = pytesseract.image_to_string(img)
                        if page_ocr and page_ocr.strip():
                            ocr_parts.append(page_ocr.strip())

                    if ocr_parts:
                        text_parts = ocr_parts
                except Exception:
                    # Gracefully handle missing Tesseract binary, poppler error, or OCR failure
                    pass

        elif ext == '.docx':
            if hasattr(file_obj_or_path, 'read'):
                file_obj_or_path.seek(0)
                stream = io.BytesIO(file_obj_or_path.read())
                file_obj_or_path.seek(0)
                doc = docx.Document(stream)
            else:
                doc = docx.Document(file_obj_or_path)

            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)

    except Exception:
        # If extraction fails on corrupted or malformed files, gracefully return empty text
        return ''

    return '\n'.join(text_parts).strip()



def guess_document_type(text):
    """
    Guess the document type from keywords in the extracted text.
    Returns one of Document.DOCUMENT_TYPE_CHOICES (defaulting to 'other').
    """
    if not text:
        return Document.DOC_TYPE_OTHER

    lower_text = text.lower()

    # Define keyword patterns for each document type with priority ordering
    patterns = [
        (
            Document.DOC_TYPE_SAFETY_CIRCULAR,
            [
                r'\bsafety\s+circular\b',
                r'\bsafety\s+notice\b',
                r'\bsafety\s+advisory\b',
                r'\bsafety\s+instruction\b',
                r'\bhazard\s+alert\b',
                r'\bemergency\s+protocol\b',
                r'\bsafety\b',
            ],
        ),
        (
            Document.DOC_TYPE_TENDER,
            [
                r'\bnotice\s+inviting\s+tender\b',
                r'\be-tender\b',
                r'\btender\s+notice\b',
                r'\btender\s+document\b',
                r'\btender\b',
                r'\brequest\s+for\s+proposal\b',
                r'\brfp\b',
                r'\bbidding\s+document\b',
                r'\bprocurement\b',
            ],
        ),
        (
            Document.DOC_TYPE_INVOICE,
            [
                r'\btax\s+invoice\b',
                r'\bproforma\s+invoice\b',
                r'\binvoice\b',
                r'\bbill\s+to\b',
                r'\bpayment\s+receipt\b',
                r'\btotal\s+amount\s+due\b',
                r'\bremittance\b',
            ],
        ),
        (
            Document.DOC_TYPE_LEGAL,
            [
                r'\blegal\s+notice\b',
                r'\bcourt\s+of\b',
                r'\baffidavit\b',
                r'\barbitration\b',
                r'\bstatutory\s+notice\b',
                r'\blegal\s+counsel\b',
                r'\bjurisdiction\b',
                r'\bpetition\b',
            ],
        ),
        (
            Document.DOC_TYPE_CONTRACT,
            [
                r'\bcontract\s+agreement\b',
                r'\bmemorandum\s+of\s+understanding\b',
                r'\bmou\b',
                r'\bcontract\b',
                r'\bagreement\b',
                r'\bnon-disclosure\s+agreement\b',
                r'\bnda\b',
                r'\bservice\s+level\s+agreement\b',
                r'\bsla\b',
            ],
        ),
        (
            Document.DOC_TYPE_REPORT,
            [
                r'\bannual\s+report\b',
                r'\baudit\s+report\b',
                r'\binspection\s+report\b',
                r'\bprogress\s+report\b',
                r'\bmonthly\s+report\b',
                r'\bquarterly\s+report\b',
                r'\bfeasibility\s+report\b',
                r'\bstatus\s+report\b',
                r'\breport\b',
            ],
        ),
    ]

    for doc_type, regex_list in patterns:
        for regex in regex_list:
            if re.search(regex, lower_text):
                return doc_type

    return Document.DOC_TYPE_OTHER
