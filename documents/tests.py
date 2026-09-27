import io
from unittest.mock import patch
from django.conf import settings
from django.contrib import admin
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
import docx

from accounts.models import Department
from accounts.permissions import (
    CONFIDENTIALITY_INTERNAL,
    CONFIDENTIALITY_PUBLIC,
    CONFIDENTIALITY_RESTRICTED,
)
from documents.models import Document, DocumentVersion
from documents.utils import extract_text, guess_document_type

User = get_user_model()


class DocumentModelsTest(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name='Operations & Station Management')
        self.user = User.objects.create_user(
            username='test_officer',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.dept,
        )

    def test_create_document_and_version(self):
        doc = Document.objects.create(
            title='Station Safety Protocol 2026',
            document_type=Document.DOC_TYPE_SAFETY_CIRCULAR,
            department=self.dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_UPLOADED,
            uploaded_by=self.user,
        )
        self.assertEqual(str(doc), 'Station Safety Protocol 2026')
        self.assertIsNone(doc.current_version)
        self.assertEqual(doc.status, Document.STATUS_UPLOADED)
        self.assertEqual(doc.confidentiality, CONFIDENTIALITY_INTERNAL)

        dummy_file = SimpleUploadedFile(
            'safety_protocol_v1.pdf',
            b'%PDF-1.4 dummy pdf content',
            content_type='application/pdf',
        )
        version = DocumentVersion.objects.create(
            document=doc,
            version_number=1,
            file=dummy_file,
            extracted_text='Dummy extracted text',
            text_found=True,
            uploaded_by=self.user,
            change_note='Initial draft upload',
        )
        doc.current_version = version
        doc.save()

        self.assertEqual(doc.current_version, version)
        self.assertEqual(str(version), 'Station Safety Protocol 2026 (v1)')
        self.assertEqual(version.version_number, 1)
        self.assertTrue(version.text_found)
        self.assertEqual(version.change_note, 'Initial draft upload')
        self.assertEqual(doc.versions.count(), 1)

    def test_document_confidentiality_choices(self):
        for choice in [CONFIDENTIALITY_PUBLIC, CONFIDENTIALITY_INTERNAL, CONFIDENTIALITY_RESTRICTED]:
            doc = Document.objects.create(
                title=f'Document {choice}',
                document_type=Document.DOC_TYPE_REPORT,
                department=self.dept,
                confidentiality=choice,
            )
            self.assertEqual(doc.confidentiality, choice)

    def test_admin_registration(self):
        self.assertIn(Document, admin.site._registry)
        self.assertIn(DocumentVersion, admin.site._registry)

    def test_media_root_setting(self):
        self.assertEqual(settings.MEDIA_ROOT, settings.BASE_DIR / 'protected_files')


class DocumentHelpersTest(TestCase):
    def test_guess_document_type(self):
        self.assertEqual(guess_document_type('This is a Notice Inviting Tender for rolling stock'), Document.DOC_TYPE_TENDER)
        self.assertEqual(guess_document_type('Contract Agreement and Terms between parties'), Document.DOC_TYPE_CONTRACT)
        self.assertEqual(guess_document_type('Tax Invoice and Bill To Kochi Metro Rail Ltd'), Document.DOC_TYPE_INVOICE)
        self.assertEqual(guess_document_type('Safety Circular: Emergency brake protocol'), Document.DOC_TYPE_SAFETY_CIRCULAR)
        self.assertEqual(guess_document_type('Legal notice from arbitration high court'), Document.DOC_TYPE_LEGAL)
        self.assertEqual(guess_document_type('Monthly Progress Report and Inspection Audit'), Document.DOC_TYPE_REPORT)
        self.assertEqual(guess_document_type('Random text with no matching keywords'), Document.DOC_TYPE_OTHER)
        self.assertEqual(guess_document_type(''), Document.DOC_TYPE_OTHER)

    def test_extract_text_from_docx(self):
        doc = docx.Document()
        doc.add_paragraph('Test safety circular content for extraction.')
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        uploaded_docx = SimpleUploadedFile('test.docx', doc_io.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        text = extract_text(uploaded_docx)
        self.assertIn('Test safety circular content', text)

    def test_extract_text_on_corrupt_file(self):
        corrupt_file = SimpleUploadedFile('corrupt.pdf', b'not a real pdf content', content_type='application/pdf')
        text = extract_text(corrupt_file)
        self.assertEqual(text, '')


class DocumentUploadViewTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.employee = User.objects.create_user(
            username='emp_user',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_user',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.admin = User.objects.create_superuser(
            username='admin_user',
            password='password123',
            role=User.ROLE_ADMIN,
            department=self.fin_dept,
        )

    def test_unauthenticated_redirect(self):
        response = self.client.get(reverse('document_upload'))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response.url)

    def test_upload_page_get_as_employee(self):
        self.client.login(username='emp_user', password='password123')
        response = self.client.get(reverse('document_upload'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'documents/upload.html')
        # Confidentiality choices should NOT have restricted for employee
        form = response.context['form']
        conf_choices = [c[0] for c in form.fields['confidentiality'].choices]
        self.assertIn(CONFIDENTIALITY_PUBLIC, conf_choices)
        self.assertIn(CONFIDENTIALITY_INTERNAL, conf_choices)
        self.assertNotIn(CONFIDENTIALITY_RESTRICTED, conf_choices)

    def test_upload_page_get_as_officer(self):
        self.client.login(username='off_user', password='password123')
        response = self.client.get(reverse('document_upload'))
        self.assertEqual(response.status_code, 200)
        form = response.context['form']
        conf_choices = [c[0] for c in form.fields['confidentiality'].choices]
        self.assertIn(CONFIDENTIALITY_RESTRICTED, conf_choices)

    def test_successful_upload_with_auto_detect(self):
        self.client.login(username='off_user', password='password123')

        doc = docx.Document()
        doc.add_paragraph('Notice Inviting Tender for Track Maintenance Equipment')
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        file = SimpleUploadedFile('tender_notice.docx', doc_io.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        response = self.client.post(reverse('document_upload'), {
            'file': file,
            'title': '',  # Empty title -> auto-fill from filename
            'document_type': 'auto',
            'confidentiality': CONFIDENTIALITY_INTERNAL,
        })
        self.assertRedirects(response, reverse('landing'))

        doc_obj = Document.objects.filter(uploaded_by=self.officer).first()
        self.assertIsNotNone(doc_obj)
        self.assertEqual(doc_obj.title, 'tender notice')
        self.assertEqual(doc_obj.document_type, Document.DOC_TYPE_TENDER)
        self.assertEqual(doc_obj.department, self.ops_dept)
        self.assertIsNotNone(doc_obj.current_version)
        self.assertTrue(doc_obj.current_version.text_found)
        self.assertIn('Notice Inviting Tender', doc_obj.current_version.extracted_text)

    def test_admin_can_select_other_department(self):
        self.client.login(username='admin_user', password='password123')
        file = SimpleUploadedFile('report.docx', b'dummy content', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        response = self.client.post(reverse('document_upload'), {
            'file': file,
            'title': 'Operations Safety Report',
            'document_type': Document.DOC_TYPE_REPORT,
            'confidentiality': CONFIDENTIALITY_RESTRICTED,
            'department': self.ops_dept.id,
        })
        self.assertRedirects(response, reverse('landing'))

        doc_obj = Document.objects.filter(title='Operations Safety Report').first()
        self.assertEqual(doc_obj.department, self.ops_dept)

    def test_invalid_file_extension(self):
        self.client.login(username='emp_user', password='password123')
        invalid_file = SimpleUploadedFile('image.png', b'image data', content_type='image/png')
        response = self.client.post(reverse('document_upload'), {
            'file': invalid_file,
            'title': 'Test',
            'document_type': 'auto',
            'confidentiality': CONFIDENTIALITY_PUBLIC,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['form'], 'file', 'Only PDF and DOCX files are allowed.')

    def test_file_too_large(self):
        self.client.login(username='emp_user', password='password123')
        large_file = SimpleUploadedFile('large.pdf', b'x' * (10 * 1024 * 1024 + 1), content_type='application/pdf')
        response = self.client.post(reverse('document_upload'), {
            'file': large_file,
            'title': 'Large File',
            'document_type': 'auto',
            'confidentiality': CONFIDENTIALITY_PUBLIC,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['form'], 'file', 'File size cannot exceed 10 MB.')


class DocumentListViewTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.employee = User.objects.create_user(
            username='emp_user',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_user',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.admin = User.objects.create_superuser(
            username='admin_user',
            password='password123',
            role=User.ROLE_ADMIN,
            department=self.fin_dept,
        )

        # Create test documents
        self.doc_ops_pub = Document.objects.create(
            title='Ops Public Circular',
            document_type=Document.DOC_TYPE_SAFETY_CIRCULAR,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_PUBLIC,
            uploaded_by=self.employee,
        )
        self.doc_ops_int = Document.objects.create(
            title='Ops Internal Manual',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            uploaded_by=self.employee,
        )
        self.doc_ops_res = Document.objects.create(
            title='Ops Restricted Protocol',
            document_type=Document.DOC_TYPE_LEGAL,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_RESTRICTED,
            uploaded_by=self.officer,
        )
        self.doc_fin_int = Document.objects.create(
            title='Fin Internal Ledger',
            document_type=Document.DOC_TYPE_INVOICE,
            department=self.fin_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            uploaded_by=self.admin,
        )
        self.doc_fin_res = Document.objects.create(
            title='Fin Restricted Audit',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.fin_dept,
            confidentiality=CONFIDENTIALITY_RESTRICTED,
            uploaded_by=self.admin,
        )

    def test_unauthenticated_redirect(self):
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response.url)

    def test_employee_sees_only_own_department_public_and_internal(self):
        self.client.login(username='emp_user', password='password123')
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        doc_titles = [d.title for d in docs]
        self.assertIn('Ops Public Circular', doc_titles)
        self.assertIn('Ops Internal Manual', doc_titles)
        self.assertNotIn('Ops Restricted Protocol', doc_titles)
        self.assertNotIn('Fin Internal Ledger', doc_titles)
        self.assertNotIn('Fin Restricted Audit', doc_titles)

    def test_officer_sees_all_public_internal_and_own_restricted(self):
        self.client.login(username='off_user', password='password123')
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        doc_titles = [d.title for d in docs]
        self.assertIn('Ops Public Circular', doc_titles)
        self.assertIn('Ops Internal Manual', doc_titles)
        self.assertIn('Ops Restricted Protocol', doc_titles)
        self.assertIn('Fin Internal Ledger', doc_titles)
        self.assertNotIn('Fin Restricted Audit', doc_titles)

    def test_admin_sees_all_documents(self):
        self.client.login(username='admin_user', password='password123')
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 5)

    def test_empty_state_when_no_documents(self):
        Document.objects.all().delete()
        self.client.login(username='emp_user', password='password123')
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'No documents visible')


class DashboardDataUpdatedTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.employee = User.objects.create_user(
            username='emp_dash',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_dash',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )

        self.doc1 = Document.objects.create(
            title='Ops Circular 1',
            document_type=Document.DOC_TYPE_SAFETY_CIRCULAR,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_PUBLIC,
            status=Document.STATUS_PENDING_APPROVAL,
            uploaded_by=self.employee,
        )
        self.doc2 = Document.objects.create(
            title='Fin Ledger 2',
            document_type=Document.DOC_TYPE_INVOICE,
            department=self.fin_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_APPROVED,
            uploaded_by=self.officer,
        )

    def test_employee_dashboard_counts(self):
        from accounts.views import get_dashboard_data
        data = get_dashboard_data(self.employee)
        cards = {c['key']: c for c in data['cards']}
        self.assertEqual(cards['total_docs']['value'], 1)  # Only doc1 is visible to emp
        self.assertEqual(cards['total_docs']['link'], reverse('document_list'))
        self.assertEqual(cards['approvals']['label'], 'My submissions')
        self.assertEqual(cards['approvals']['value'], 1)  # emp uploaded doc1 (pending approval)
        self.assertEqual(cards['recent_uploads']['value'], 1)
        self.assertEqual(cards['notifications']['value'], 0)

    def test_officer_dashboard_counts(self):
        from accounts.views import get_dashboard_data
        data = get_dashboard_data(self.officer)
        cards = {c['key']: c for c in data['cards']}
        self.assertEqual(cards['total_docs']['value'], 2)  # doc1 (public) and doc2 (internal) are visible
        self.assertEqual(cards['approvals']['label'], 'Awaiting approval')
        self.assertEqual(cards['approvals']['value'], 1)  # doc1 in ops_dept is pending approval


class DocumentDetailAndDownloadTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.ops_employee = User.objects.create_user(
            username='emp_ops',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.ops_officer = User.objects.create_user(
            username='off_ops',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.fin_employee = User.objects.create_user(
            username='emp_fin',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.fin_dept,
        )

        self.doc_ops_internal = Document.objects.create(
            title='Operations Rulebook',
            document_type=Document.DOC_TYPE_SAFETY_CIRCULAR,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_APPROVED,
            uploaded_by=self.ops_officer,
        )
        self.dummy_file_v1 = SimpleUploadedFile('rulebook_v1.pdf', b'PDF-v1 content', content_type='application/pdf')
        self.v1 = DocumentVersion.objects.create(
            document=self.doc_ops_internal,
            version_number=1,
            file=self.dummy_file_v1,
            uploaded_by=self.ops_officer,
            change_note='First edition',
        )
        self.dummy_file_v2 = SimpleUploadedFile('rulebook_v2.pdf', b'PDF-v2 content', content_type='application/pdf')
        self.v2 = DocumentVersion.objects.create(
            document=self.doc_ops_internal,
            version_number=2,
            file=self.dummy_file_v2,
            uploaded_by=self.ops_officer,
            change_note='Second edition revised',
        )
        self.doc_ops_internal.current_version = self.v2
        self.doc_ops_internal.save()

        self.doc_ops_restricted = Document.objects.create(
            title='Confidential Ops Incident',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_RESTRICTED,
            uploaded_by=self.ops_officer,
        )
        self.dummy_file_restricted = SimpleUploadedFile('incident.pdf', b'Restricted content', content_type='application/pdf')
        self.v_restricted = DocumentVersion.objects.create(
            document=self.doc_ops_restricted,
            version_number=1,
            file=self.dummy_file_restricted,
            uploaded_by=self.ops_officer,
        )
        self.doc_ops_restricted.current_version = self.v_restricted
        self.doc_ops_restricted.save()

    def test_unauthenticated_detail_redirect(self):
        response = self.client.get(reverse('document_detail', args=[self.doc_ops_internal.id]))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response.url)

    def test_allowed_user_can_view_detail(self):
        self.client.login(username='emp_ops', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc_ops_internal.id]))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'documents/document_detail.html')
        self.assertContains(response, 'Operations Rulebook')
        self.assertContains(response, 'v1')
        self.assertContains(response, 'v2')
        self.assertContains(response, 'First edition')
        self.assertContains(response, 'Second edition revised')

    def test_disallowed_user_gets_404_for_detail(self):
        # emp_ops cannot view restricted doc in own department
        self.client.login(username='emp_ops', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc_ops_restricted.id]))
        self.assertEqual(response.status_code, 404)

        # emp_fin cannot view internal doc from ops department
        self.client.login(username='emp_fin', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc_ops_internal.id]))
        self.assertEqual(response.status_code, 404)

    def test_nonexistent_document_detail_gets_404(self):
        self.client.login(username='off_ops', password='password123')
        response = self.client.get(reverse('document_detail', args=[9999]))
        self.assertEqual(response.status_code, 404)

    def test_allowed_user_can_download_version(self):
        self.client.login(username='emp_ops', password='password123')
        response = self.client.get(reverse('document_version_download', args=[self.doc_ops_internal.id, self.v1.id]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Disposition'], f'attachment; filename="{self.v1.file.name.split("/")[-1]}"')
        content = b''.join(response.streaming_content)
        self.assertEqual(content, b'PDF-v1 content')

    def test_disallowed_user_gets_404_for_download(self):
        self.client.login(username='emp_ops', password='password123')
        response = self.client.get(reverse('document_version_download', args=[self.doc_ops_restricted.id, self.v_restricted.id]))
        self.assertEqual(response.status_code, 404)

    def test_document_list_links_to_detail(self):
        self.client.login(username='emp_ops', password='password123')
        response = self.client.get(reverse('document_list'))
        self.assertEqual(response.status_code, 200)
        detail_url = reverse('document_detail', args=[self.doc_ops_internal.id])
        self.assertContains(response, f'href="{detail_url}"')


class DocumentNewVersionTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.uploader = User.objects.create_user(
            username='emp_uploader',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.other_officer = User.objects.create_user(
            username='off_other',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.admin = User.objects.create_superuser(
            username='admin_boss',
            password='password123',
            role=User.ROLE_ADMIN,
            department=self.fin_dept,
        )

        self.doc = Document.objects.create(
            title='Operations Schedule 2026',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            uploaded_by=self.uploader,
        )
        self.v1 = DocumentVersion.objects.create(
            document=self.doc,
            version_number=1,
            file=SimpleUploadedFile('sched_v1.pdf', b'Schedule v1 content', content_type='application/pdf'),
            uploaded_by=self.uploader,
            change_note='Initial draft',
        )
        self.doc.current_version = self.v1
        self.doc.save()

    def test_uploader_sees_upload_new_version_form(self):
        self.client.login(username='emp_uploader', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['can_upload_version'])
        self.assertContains(response, 'Upload New Version')

    def test_admin_sees_upload_new_version_form(self):
        self.client.login(username='admin_boss', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['can_upload_version'])
        self.assertContains(response, 'Upload New Version')

    def test_non_uploader_does_not_see_upload_new_version_form(self):
        self.client.login(username='off_other', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['can_upload_version'])
        self.assertNotContains(response, 'Upload New Version')

    def test_uploader_can_upload_new_version(self):
        self.client.login(username='emp_uploader', password='password123')

        doc_file = docx.Document()
        doc_file.add_paragraph('Updated timetable and station timings')
        doc_io = io.BytesIO()
        doc_file.save(doc_io)
        doc_io.seek(0)
        uploaded_docx = SimpleUploadedFile('sched_v2.docx', doc_io.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        response = self.client.post(reverse('document_upload_new_version', args=[self.doc.id]), {
            'file': uploaded_docx,
            'change_note': 'Added weekend timetable',
        })
        self.assertRedirects(response, reverse('document_detail', args=[self.doc.id]))

        self.doc.refresh_from_db()
        self.assertEqual(self.doc.versions.count(), 2)
        v2 = self.doc.versions.get(version_number=2)
        self.assertEqual(self.doc.current_version, v2)
        self.assertEqual(v2.change_note, 'Added weekend timetable')
        self.assertTrue(v2.text_found)
        self.assertIn('Updated timetable', v2.extracted_text)

        # Confirm v1 is still present and untouched
        v1 = self.doc.versions.get(version_number=1)
        self.assertEqual(v1.change_note, 'Initial draft')

    def test_non_uploader_cannot_post_new_version(self):
        self.client.login(username='off_other', password='password123')
        uploaded_pdf = SimpleUploadedFile('sched_v2.pdf', b'New content', content_type='application/pdf')
        response = self.client.post(reverse('document_upload_new_version', args=[self.doc.id]), {
            'file': uploaded_pdf,
            'change_note': 'Unauthorized update attempt',
        })
        self.assertEqual(response.status_code, 404)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.versions.count(), 1)


class DocumentArchiveTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.uploader = User.objects.create_user(
            username='emp_uploader',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.other_employee = User.objects.create_user(
            username='emp_other',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_ops',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.admin = User.objects.create_superuser(
            username='admin_boss',
            password='password123',
            role=User.ROLE_ADMIN,
            department=self.fin_dept,
        )

        self.doc = Document.objects.create(
            title='Operations SOP 2026',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_APPROVED,
            uploaded_by=self.uploader,
        )
        self.v1 = DocumentVersion.objects.create(
            document=self.doc,
            version_number=1,
            file=SimpleUploadedFile('sop_v1.pdf', b'SOP v1 content', content_type='application/pdf'),
            uploaded_by=self.uploader,
            change_note='Initial draft',
        )
        self.doc.current_version = self.v1
        self.doc.save()

    def test_uploader_sees_archive_button(self):
        self.client.login(username='emp_uploader', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Archive')
        self.assertNotContains(response, 'Unarchive')

    def test_admin_sees_archive_button(self):
        self.client.login(username='admin_boss', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Archive')
        self.assertNotContains(response, 'Unarchive')

    def test_other_user_does_not_see_archive_button(self):
        self.client.login(username='emp_other', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Archive')
        self.assertNotContains(response, 'Unarchive')

    def test_archive_document_flow_and_status_restoration(self):
        self.client.login(username='emp_uploader', password='password123')

        # 1. Archive document
        response = self.client.post(reverse('document_toggle_archive', args=[self.doc.id]))
        self.assertRedirects(response, reverse('document_detail', args=[self.doc.id]))

        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, Document.STATUS_ARCHIVED)
        self.assertEqual(self.doc.previous_status, Document.STATUS_APPROVED)

        # 2. Detail view shows Unarchive button & Archived badge
        detail_res = self.client.get(reverse('document_detail', args=[self.doc.id]))
        self.assertEqual(detail_res.status_code, 200)
        self.assertContains(detail_res, 'Unarchive')
        self.assertContains(detail_res, 'Archived')

        # 3. Document list still displays archived document with Archived badge
        list_res = self.client.get(reverse('document_list'))
        self.assertEqual(list_res.status_code, 200)
        self.assertContains(list_res, 'Operations SOP 2026')
        self.assertContains(list_res, 'Archived')

        # 4. Unarchive document restores previous status (approved)
        unarchive_res = self.client.post(reverse('document_toggle_archive', args=[self.doc.id]))
        self.assertRedirects(unarchive_res, reverse('document_detail', args=[self.doc.id]))

        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, Document.STATUS_APPROVED)
        self.assertEqual(self.doc.previous_status, '')

    def test_unauthorized_user_cannot_toggle_archive(self):
        self.client.login(username='emp_other', password='password123')
        response = self.client.post(reverse('document_toggle_archive', args=[self.doc.id]))
        self.assertEqual(response.status_code, 404)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, Document.STATUS_APPROVED)

    def test_archived_document_excluded_from_recent_uploads(self):
        from accounts.views import get_dashboard_data

        # Prior to archiving, recent_uploads includes doc
        data_before = get_dashboard_data(self.uploader)
        cards_before = {c['key']: c for c in data_before['cards']}
        self.assertEqual(cards_before['total_docs']['value'], 1)
        self.assertEqual(cards_before['recent_uploads']['value'], 1)
        self.assertEqual(len(data_before['recent_uploads']), 1)

        # Archive the doc
        self.doc.status = Document.STATUS_ARCHIVED
        self.doc.previous_status = Document.STATUS_APPROVED
        self.doc.save()

        # After archiving: total_docs remains 1, recent_uploads becomes 0
        data_after = get_dashboard_data(self.uploader)
        cards_after = {c['key']: c for c in data_after['cards']}
        self.assertEqual(cards_after['total_docs']['value'], 1)
        self.assertEqual(cards_after['recent_uploads']['value'], 0)
        self.assertEqual(len(data_after['recent_uploads']), 0)


class DocumentSearchAndFilterTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.employee = User.objects.create_user(
            username='emp_searcher',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_searcher',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )
        self.admin = User.objects.create_superuser(
            username='admin_searcher',
            password='password123',
            role=User.ROLE_ADMIN,
            department=self.fin_dept,
        )

        # Document 1: Ops Tender, Public, Approved
        self.doc1 = Document.objects.create(
            title='Track Maintenance Tender 2026',
            document_type=Document.DOC_TYPE_TENDER,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_PUBLIC,
            status=Document.STATUS_APPROVED,
            uploaded_by=self.employee,
        )
        self.v1_doc1 = DocumentVersion.objects.create(
            document=self.doc1,
            version_number=1,
            file=SimpleUploadedFile('tender.pdf', b'Tender content', content_type='application/pdf'),
            extracted_text='Procurement of rolling stock and braking mechanism',
            text_found=True,
            uploaded_by=self.employee,
        )
        self.doc1.current_version = self.v1_doc1
        self.doc1.save()

        # Document 2: Ops Safety Circular, Internal, Uploaded
        self.doc2 = Document.objects.create(
            title='Emergency Evacuation Protocol',
            document_type=Document.DOC_TYPE_SAFETY_CIRCULAR,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_UPLOADED,
            uploaded_by=self.employee,
        )
        self.v1_doc2 = DocumentVersion.objects.create(
            document=self.doc2,
            version_number=1,
            file=SimpleUploadedFile('evacuation.docx', b'Docx content', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
            extracted_text='Platform safety fire drill and passenger evacuation guide',
            text_found=True,
            uploaded_by=self.employee,
        )
        self.doc2.current_version = self.v1_doc2
        self.doc2.save()

        # Document 3: Finance Invoice, Internal, Pending Approval
        self.doc3 = Document.objects.create(
            title='Siemens Signaling Invoice',
            document_type=Document.DOC_TYPE_INVOICE,
            department=self.fin_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_PENDING_APPROVAL,
            uploaded_by=self.admin,
        )
        self.v1_doc3 = DocumentVersion.objects.create(
            document=self.doc3,
            version_number=1,
            file=SimpleUploadedFile('invoice.pdf', b'Invoice content', content_type='application/pdf'),
            extracted_text='Invoice number INV-9988 for signaling software license',
            text_found=True,
            uploaded_by=self.admin,
        )
        self.doc3.current_version = self.v1_doc3
        self.doc3.save()

        # Document 4: Finance Legal, Restricted, Approved
        self.doc4 = Document.objects.create(
            title='Arbitration Dispute Settlement',
            document_type=Document.DOC_TYPE_LEGAL,
            department=self.fin_dept,
            confidentiality=CONFIDENTIALITY_RESTRICTED,
            status=Document.STATUS_APPROVED,
            uploaded_by=self.admin,
        )
        self.v1_doc4 = DocumentVersion.objects.create(
            document=self.doc4,
            version_number=1,
            file=SimpleUploadedFile('legal.pdf', b'Legal content', content_type='application/pdf'),
            extracted_text='High court commercial arbitration hearing minutes',
            text_found=True,
            uploaded_by=self.admin,
        )
        self.doc4.current_version = self.v1_doc4
        self.doc4.save()

    def test_search_by_title_case_insensitive(self):
        self.client.login(username='off_searcher', password='password123')
        # Search 'TRACK' matching 'Track Maintenance Tender 2026'
        response = self.client.get(reverse('document_list'), {'q': 'TRACK'})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc1)
        self.assertContains(response, 'Track Maintenance Tender 2026')
        self.assertNotContains(response, 'Emergency Evacuation Protocol')

    def test_search_by_extracted_text_case_insensitive(self):
        self.client.login(username='off_searcher', password='password123')
        # Search 'braking mechanism' which is inside doc1 version extracted_text
        response = self.client.get(reverse('document_list'), {'q': 'BRAKING MECHANISM'})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc1)

    def test_filter_by_document_type(self):
        self.client.login(username='off_searcher', password='password123')
        response = self.client.get(reverse('document_list'), {'document_type': Document.DOC_TYPE_SAFETY_CIRCULAR})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc2)

    def test_filter_by_department(self):
        self.client.login(username='off_searcher', password='password123')
        # Officer sees Public & Internal across departments: doc1 (Ops), doc2 (Ops), doc3 (Fin)
        response = self.client.get(reverse('document_list'), {'department': self.fin_dept.id})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc3)

    def test_filter_by_status(self):
        self.client.login(username='off_searcher', password='password123')
        response = self.client.get(reverse('document_list'), {'status': Document.STATUS_PENDING_APPROVAL})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc3)

    def test_filter_by_date_range(self):
        self.client.login(username='off_searcher', password='password123')
        today = self.doc1.created_date.strftime('%Y-%m-%d')
        response = self.client.get(reverse('document_list'), {
            'date_from': today,
            'date_to': today,
        })
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 3)  # doc1, doc2, doc3 created today and visible to officer

    def test_combined_search_and_filters(self):
        self.client.login(username='admin_searcher', password='password123')
        response = self.client.get(reverse('document_list'), {
            'q': 'arbitration',
            'document_type': Document.DOC_TYPE_LEGAL,
            'department': self.fin_dept.id,
            'status': Document.STATUS_APPROVED,
        })
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0], self.doc4)

    def test_permissions_respected_under_search(self):
        # Employee cannot see Finance restricted document (doc4) even if searching for its title
        self.client.login(username='emp_searcher', password='password123')
        response = self.client.get(reverse('document_list'), {'q': 'Arbitration'})
        self.assertEqual(response.status_code, 200)
        docs = response.context['documents']
        self.assertEqual(len(docs), 0)
        self.assertNotContains(response, 'Arbitration Dispute Settlement')

    def test_empty_state_with_active_filter_vs_no_documents(self):
        self.client.login(username='emp_searcher', password='password123')

        # Filter yielding no results shows search-specific empty state
        response = self.client.get(reverse('document_list'), {'q': 'nonexistent query'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'No matching documents')
        self.assertNotContains(response, 'No documents visible')

        # Clear repository to test generic no-documents empty state
        Document.objects.all().delete()
        empty_res = self.client.get(reverse('document_list'))
        self.assertEqual(empty_res.status_code, 200)
        self.assertContains(empty_res, 'No documents visible')
        self.assertNotContains(empty_res, 'No matching documents')

    def test_form_retains_selected_filter_values(self):
        self.client.login(username='admin_searcher', password='password123')
        response = self.client.get(reverse('document_list'), {
            'q': 'Signaling',
            'document_type': Document.DOC_TYPE_INVOICE,
            'department': self.fin_dept.id,
            'status': Document.STATUS_PENDING_APPROVAL,
            'date_from': '2026-01-01',
            'date_to': '2026-12-31',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'value="Signaling"')
        self.assertContains(response, 'value="2026-01-01"')
        self.assertContains(response, 'value="2026-12-31"')
        self.assertContains(response, 'selected')


class DocumentResubmitTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.ops_dept = Department.objects.create(name='Operations & Station Management')
        self.fin_dept = Department.objects.create(name='Finance & Accounts')

        self.uploader = User.objects.create_user(
            username='emp_uploader',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.other_employee = User.objects.create_user(
            username='emp_other',
            password='password123',
            role=User.ROLE_EMPLOYEE,
            department=self.ops_dept,
        )
        self.officer = User.objects.create_user(
            username='off_ops',
            password='password123',
            role=User.ROLE_OFFICER,
            department=self.ops_dept,
        )

        self.rejected_doc = Document.objects.create(
            title='Rejected Station Protocol',
            document_type=Document.DOC_TYPE_REPORT,
            department=self.ops_dept,
            confidentiality=CONFIDENTIALITY_INTERNAL,
            status=Document.STATUS_REJECTED,
            uploaded_by=self.uploader,
        )
        self.v1 = DocumentVersion.objects.create(
            document=self.rejected_doc,
            version_number=1,
            file=SimpleUploadedFile('protocol_v1.pdf', b'Protocol v1 content', content_type='application/pdf'),
            uploaded_by=self.uploader,
            change_note='Initial draft',
        )
        self.rejected_doc.current_version = self.v1
        self.rejected_doc.save()

    def test_uploader_sees_resubmit_form_when_rejected(self):
        self.client.login(username='emp_uploader', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.rejected_doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['can_resubmit'])
        self.assertContains(response, 'Resubmit Document for Approval')

    def test_uploader_does_not_see_resubmit_form_when_not_rejected(self):
        self.rejected_doc.status = Document.STATUS_APPROVED
        self.rejected_doc.save()

        self.client.login(username='emp_uploader', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.rejected_doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['can_resubmit'])
        self.assertNotContains(response, 'Resubmit Document for Approval')

    def test_non_uploader_does_not_see_resubmit_form(self):
        self.client.login(username='off_ops', password='password123')
        response = self.client.get(reverse('document_detail', args=[self.rejected_doc.id]))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['can_resubmit'])
        self.assertNotContains(response, 'Resubmit Document for Approval')

    def test_successful_resubmit_without_file(self):
        from audit.models import AuditLog
        from workflow.models import Approval

        self.client.login(username='emp_uploader', password='password123')
        response = self.client.post(reverse('document_resubmit', args=[self.rejected_doc.id]), {
            'title': 'Revised Station Protocol 2026',
            'document_type': Document.DOC_TYPE_SAFETY_CIRCULAR,
            'confidentiality': CONFIDENTIALITY_PUBLIC,
            'note': 'Fixed emergency escalation contacts as requested.',
        })
        self.assertRedirects(response, reverse('document_detail', args=[self.rejected_doc.id]))

        self.rejected_doc.refresh_from_db()
        self.assertEqual(self.rejected_doc.title, 'Revised Station Protocol 2026')
        self.assertEqual(self.rejected_doc.document_type, Document.DOC_TYPE_SAFETY_CIRCULAR)
        self.assertEqual(self.rejected_doc.confidentiality, CONFIDENTIALITY_PUBLIC)
        self.assertEqual(self.rejected_doc.status, Document.STATUS_PENDING_APPROVAL)
        # Version remains v1
        self.assertEqual(self.rejected_doc.current_version, self.v1)
        self.assertEqual(self.rejected_doc.versions.count(), 1)

        # Confirm new Approval record in workflow
        new_approval = Approval.objects.filter(document=self.rejected_doc, status=Approval.STATUS_PENDING).first()
        self.assertIsNotNone(new_approval)
        self.assertEqual(new_approval.comment, 'Fixed emergency escalation contacts as requested.')

        # Confirm audit log entry
        audit_entry = AuditLog.objects.filter(document=self.rejected_doc, action='resubmit').first()
        self.assertIsNotNone(audit_entry)
        self.assertEqual(audit_entry.user, self.uploader)
        self.assertIn('New file attached: No', audit_entry.details)
        self.assertIn('Fixed emergency escalation contacts', audit_entry.details)

    def test_successful_resubmit_with_new_file(self):
        import docx
        import io
        from audit.models import AuditLog
        from workflow.models import Approval

        doc_file = docx.Document()
        doc_file.add_paragraph('Updated emergency protocol text and phone tree')
        doc_io = io.BytesIO()
        doc_file.save(doc_io)
        doc_io.seek(0)
        uploaded_docx = SimpleUploadedFile('protocol_v2.docx', doc_io.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        self.client.login(username='emp_uploader', password='password123')
        response = self.client.post(reverse('document_resubmit', args=[self.rejected_doc.id]), {
            'file': uploaded_docx,
            'title': 'Revised Protocol With New File',
            'document_type': Document.DOC_TYPE_SAFETY_CIRCULAR,
            'confidentiality': CONFIDENTIALITY_INTERNAL,
            'note': 'Replaced file with revised edition.',
        })
        self.assertRedirects(response, reverse('document_detail', args=[self.rejected_doc.id]))

        self.rejected_doc.refresh_from_db()
        self.assertEqual(self.rejected_doc.title, 'Revised Protocol With New File')
        self.assertEqual(self.rejected_doc.status, Document.STATUS_PENDING_APPROVAL)
        self.assertEqual(self.rejected_doc.versions.count(), 2)

        v2 = self.rejected_doc.versions.get(version_number=2)
        self.assertEqual(self.rejected_doc.current_version, v2)
        self.assertTrue(v2.text_found)
        self.assertIn('Updated emergency protocol text', v2.extracted_text)
        self.assertEqual(v2.change_note, 'Replaced file with revised edition.')

        # Confirm approval record
        new_approval = Approval.objects.filter(document=self.rejected_doc, status=Approval.STATUS_PENDING).first()
        self.assertIsNotNone(new_approval)
        self.assertEqual(new_approval.comment, 'Replaced file with revised edition.')

        # Confirm audit log
        audit_entry = AuditLog.objects.filter(document=self.rejected_doc, action='resubmit').first()
        self.assertIsNotNone(audit_entry)
        self.assertIn('New file attached: Yes', audit_entry.details)
        self.assertIn('Replaced file with revised edition.', audit_entry.details)

    def test_unauthorized_resubmit_attempt(self):
        self.client.login(username='emp_other', password='password123')
        response = self.client.post(reverse('document_resubmit', args=[self.rejected_doc.id]), {
            'title': 'Hacked Title',
            'document_type': Document.DOC_TYPE_SAFETY_CIRCULAR,
            'confidentiality': CONFIDENTIALITY_INTERNAL,
        })
        self.assertEqual(response.status_code, 404)
        self.rejected_doc.refresh_from_db()
        self.assertEqual(self.rejected_doc.status, Document.STATUS_REJECTED)







